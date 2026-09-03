"""
build_tp.py — Convierte solucion.md a .docx con carátula como primera hoja.
Uso: python tps/build_tp.py <numero>
     python tps/build_tp.py 1
"""

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tools.docx_builder import DocxBuilder
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

LOGO_PATH = Path(__file__).resolve().parents[1] / "tools" / "logo_unpsjb.png"

MATERIA = "Administración de Redes y Seguridad"
JTP = "Lucas Krmpotic"
ALUMNO = "Gonzalo Agú"


def extract_title(consigna_path: Path, tp_number: int) -> str:
    if consigna_path.exists():
        first_line = consigna_path.read_text(encoding="utf-8").split("\n")[0]
        m = re.match(r"^#\s+(.*)", first_line)
        if m:
            return m.group(1).strip()
    return f"Trabajo Práctico {tp_number}"


def _empty_paragraph(doc, size=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    pf = p.paragraph_format
    pf.line_spacing = Pt(size) if size else Pt(2)
    return p


def _centered_run(doc, text, *, font_size, bold=False, color, font_name="Calibri"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    run.font.color.rgb = color
    return p


def _centered_image(doc, path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    return p


def add_cover(doc, titulo: str):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    _empty_paragraph(doc, 14)

    if LOGO_PATH.exists():
        _centered_image(doc, LOGO_PATH, width=Cm(2.8))
        _empty_paragraph(doc, 10)

    _centered_run(doc, "UNIVERSIDAD NACIONAL DE LA PATAGONIA SAN JUAN BOSCO",
                  font_size=Pt(10), bold=True, color=RGBColor(0x44, 0x44, 0x44))
    _empty_paragraph(doc, 4)
    _centered_run(doc, "Facultad de Ingeniería",
                  font_size=Pt(9), bold=False, color=RGBColor(0x66, 0x66, 0x66))
    _centered_run(doc, "Ingeniería en Sistemas de Información",
                  font_size=Pt(9), bold=False, color=RGBColor(0x66, 0x66, 0x66))

    _empty_paragraph(doc, 8)

    # Separador horizontal fino
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(0)
    run_line = p_line.add_run("━" * 50)
    run_line.font.size = Pt(8)
    run_line.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    _empty_paragraph(doc, 8)

    _centered_run(doc, titulo.upper(),
                  font_size=Pt(26), bold=True, color=RGBColor(0x1A, 0x1A, 0x2E))

    _empty_paragraph(doc, 6)

    _centered_run(doc, MATERIA,
                  font_size=Pt(14), bold=False, color=RGBColor(0x33, 0x33, 0x33))

    _empty_paragraph(doc, 20)

    # Bloque de info
    for label, value in [
        ("Jefe de Trabajo Práctico", JTP),
        ("Alumno", ALUMNO),
    ]:
        _centered_run(doc, label,
                      font_size=Pt(10), bold=False, color=RGBColor(0x77, 0x77, 0x77))
        _centered_run(doc, value,
                      font_size=Pt(13), bold=True, color=RGBColor(0x1A, 0x1A, 0x2E))
        _empty_paragraph(doc, 10)

    doc.add_page_break()


# ── Parser de Markdown simple ──────────────────────────────────────────────

def add_md_content(builder, md_text: str):
    """Parsea markdown básico y agrega contenido al builder."""
    doc = builder.doc
    lines = md_text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.strip().startswith("```"):
            lang = line.strip().lstrip("`").strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1

            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # Table
        if "|" in line and line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1

            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                rows = []
                for tl in table_lines[2:]:
                    row = [c.strip() for c in tl.split("|")[1:-1]]
                    rows.append(row)
                builder.add_table(headers, rows)
            continue

        # Heading
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            builder.add_heading(text, level=level)
            i += 1
            continue

        # Blockquote
        if line.strip().startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_text = re.sub(r"^>\s?", "", lines[i].strip())
                quote_lines.append(quote_text)
                i += 1
            full_quote = " ".join(quote_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            p.paragraph_format.right_indent = Cm(1.5)
            run = p.add_run(f'"{full_quote}"')
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        # Horizontal rule
        if re.match(r"^---+\s*$", line):
            builder.add_separator()
            i += 1
            continue

        # Numbered list
        m_num = re.match(r"^(\d+)\.\s+(.*)", line)
        if m_num:
            builder.add_numbered(m_num.group(2).strip())
            i += 1
            continue

        # Bullet list
        m_bullet = re.match(r"^[-*]\s+(.*)", line)
        if m_bullet:
            builder.add_bullet(m_bullet.group(1).strip())
            i += 1
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Regular paragraph
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith("#") \
                and not lines[i].strip().startswith("|") and not lines[i].strip().startswith(">") \
                and not lines[i].strip().startswith("```") and not lines[i].strip().startswith("-") \
                and not lines[i].strip().startswith("*") and not re.match(r"^\d+\.\s", lines[i]) \
                and not re.match(r"^---", lines[i]):
            para_lines.append(lines[i])
            i += 1

        text = " ".join(para_lines)
        parts = re.split(r"\*\*(.*?)\*\*", text)
        if len(parts) > 1:
            p = doc.add_paragraph()
            for idx, part in enumerate(parts):
                if not part:
                    continue
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.bold = True
        else:
            builder.add_paragraph(text)


def find_solucion(tp_dir: Path) -> Path | None:
    """Busca solucion.md dentro de la carpeta del TP."""
    # Buscar en tps/N/A/solucion.md, tps/N/B/solucion.md, etc.
    for subdir in sorted(tp_dir.iterdir()):
        if subdir.is_dir():
            candidate = subdir / "solucion.md"
            if candidate.exists():
                return candidate
    # Buscar directamente en tps/N/solucion.md
    candidate = tp_dir / "solucion.md"
    if candidate.exists():
        return candidate
    return None


def build_tp(tp_number: int):
    """Construye el .docx completo: carátula + contenido del markdown."""
    tps_root = Path(__file__).resolve().parent
    tp_dir = tps_root / str(tp_number)

    if not tp_dir.exists():
        print(f"Error: no existe tps/{tp_number}/")
        sys.exit(1)

    consigna = tp_dir / "consigna.md"
    titulo = extract_title(consigna, tp_number)

    solucion = find_solucion(tp_dir)
    if not solucion:
        print(f"Error: no se encontró solucion.md en tps/{tp_number}/")
        sys.exit(1)

    md_content = solucion.read_text(encoding="utf-8")

    builder = DocxBuilder()
    add_cover(builder.doc, titulo)
    add_md_content(builder, md_content)

    output = tp_dir / "trabajo_practico.docx"
    result = builder.save(output)
    print(f"Documento creado: {result}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python tps/build_tp.py <numero_tp>")
        print("Ejemplo: python tps/build_tp.py 1")
        sys.exit(1)

    try:
        tp_number = int(sys.argv[1])
    except ValueError:
        print(f"Error: '{sys.argv[1]}' no es un número válido")
        sys.exit(1)

    build_tp(tp_number)
